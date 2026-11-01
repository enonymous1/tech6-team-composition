"""Reformat a qualifications summary matrix (XLSX) with rich-text bolding.

The bold phrases come from a reference Word document: runs marked bold
in the document's first table are re-applied as rich-text bold segments
in the matching matrix cells. v1 hardcoded absolute paths on the
original machine for every input and output, so it only ran there;
every path here is an explicit argument.

Input cells use literal ``<br>`` separators, which become newlines, and
bullet columns get a ``•`` prefix per line. Formulas are carried over
as formulas.
"""

from __future__ import annotations

import datetime
import logging
import math
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl
import xlsxwriter
from docx import Document

logger = logging.getLogger(__name__)

_CONTROL_CHARS_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_FONT_NAME = "Gill Sans MT"
_FONT_SIZE = 10


@dataclass(frozen=True)
class MatrixCell:
    raw_value: Any
    processed_text: str
    is_header: bool
    is_center: bool
    is_bullet: bool
    keywords: tuple[str, ...]
    col_index: int
    row_index: int


def get_bold_map(docx_path: str | Path) -> dict[tuple[str, int], list[str]]:
    """Extract bold phrases from the first table of a Word document.

    Returns ``{(expert name, column index): [bold phrases]}``.
    """
    path = Path(docx_path)
    if not path.exists():
        raise FileNotFoundError(f"Reference document not found: {path}")
    doc = Document(str(path))
    bold_map: dict[tuple[str, int], list[str]] = {}
    if not doc.tables:
        return bold_map
    table = doc.tables[0]
    for r_idx, row in enumerate(table.rows):
        if r_idx == 0:
            continue  # Header row.
        if len(row.cells) < 2:
            raise ValueError(
                f"Reference document table row {r_idx} has fewer than 2 columns; "
                "the expert name is expected in the second column"
            )
        name = row.cells[1].text.strip()
        if not name:
            continue
        for c_idx, cell in enumerate(row.cells):
            bold_phrases = [
                run.text.strip()
                for paragraph in cell.paragraphs
                for run in paragraph.runs
                if run.bold and run.text.strip()
            ]
            if bold_phrases:
                bold_map[(name, c_idx)] = bold_phrases
    return bold_map


def _classify_headers(header_row: list[str]) -> tuple[list[int], list[int]]:
    bullet_cols: list[int] = []
    center_cols: list[int] = []
    for idx, header in enumerate(header_row):
        h_lower = header.lower()
        if any(word in h_lower for word in ["education", "experience", "role", "qualification", "value"]):
            bullet_cols.append(idx)
        if any(word in h_lower for word in ["position", "name", "total", "regional"]):
            center_cols.append(idx)
    return bullet_cols, center_cols


def extract_matrix_data(
    source_path: str | Path,
    bold_map: dict[tuple[str, int], list[str]],
    sheet_name: str | None = None,
) -> list[list[MatrixCell]]:
    """Read a source matrix workbook into structured rows."""
    path = Path(source_path)
    if not path.exists():
        raise FileNotFoundError(f"Source workbook not found: {path}")
    workbook = openpyxl.load_workbook(path, data_only=False)
    try:
        ws = workbook[sheet_name] if sheet_name else workbook.active
    except KeyError as exc:
        raise ValueError(f"Worksheet {sheet_name!r} not found in {path}") from exc
    assert ws is not None
    max_row = ws.max_row
    max_col = ws.max_column

    header_row = [str(ws.cell(row=1, column=c).value or "") for c in range(1, max_col + 1)]
    bullet_cols, center_cols = _classify_headers(header_row)

    data_rows: list[list[MatrixCell]] = []
    for r in range(1, max_row + 1):
        expert_name = str(ws.cell(row=r, column=2).value or "").strip()
        row_data: list[MatrixCell] = []
        for c in range(1, max_col + 1):
            cell = ws.cell(row=r, column=c)
            val = cell.value
            # Not str(val or ""): a numeric 0 is falsy, and used to
            # collapse to a blank cell here, losing the value entirely.
            text = (str(val) if val is not None else "").replace("<br>", "\n").strip()
            text = _CONTROL_CHARS_RE.sub("", text)
            keywords: tuple[str, ...] = ()
            if r > 1:
                keywords = tuple(bold_map.get((expert_name, c - 1), []))
                if (c - 1) in bullet_cols and text:
                    lines = [line.strip() for line in text.split("\n") if line.strip()]
                    text = "\n".join(f"• {line.lstrip(' •*-')}" for line in lines)
            row_data.append(
                MatrixCell(
                    raw_value=val,
                    processed_text=text,
                    is_header=(r == 1),
                    is_center=(c - 1) in center_cols,
                    is_bullet=(c - 1) in bullet_cols,
                    keywords=keywords,
                    col_index=c - 1,
                    row_index=r - 1,
                )
            )
        data_rows.append(row_data)
    return data_rows


def write_matrix(data_rows: list[list[MatrixCell]], output_path: str | Path) -> Path:
    """Write structured matrix rows to an XlsxWriter workbook."""
    if not data_rows:
        raise ValueError("No matrix rows to write")
    out_path = Path(output_path)
    workbook = xlsxwriter.Workbook(
        str(out_path),
        {"default_format_properties": {"font_name": _FONT_NAME, "font_size": _FONT_SIZE}},
    )
    worksheet = workbook.add_worksheet("Qualifications")

    fmt_header = workbook.add_format(
        {
            "bold": True,
            "bg_color": "#BDD7EE",
            "border": 1,
            "align": "center",
            "valign": "vcenter",
            "text_wrap": True,
            "font_name": _FONT_NAME,
            "font_size": _FONT_SIZE,
        }
    )
    fmt_std = workbook.add_format(
        {
            "border": 1,
            "valign": "top",
            "text_wrap": True,
            "font_name": _FONT_NAME,
            "font_size": _FONT_SIZE,
        }
    )
    fmt_center = workbook.add_format(
        {
            "border": 1,
            "align": "center",
            "valign": "top",
            "text_wrap": True,
            "font_name": _FONT_NAME,
            "font_size": _FONT_SIZE,
        }
    )
    # Dates need a date number format, or the serial number shows as a
    # bare number when the file is read back.
    fmt_std_date = workbook.add_format(
        {
            "border": 1,
            "valign": "top",
            "text_wrap": True,
            "font_name": _FONT_NAME,
            "font_size": _FONT_SIZE,
            "num_format": "yyyy-mm-dd",
        }
    )
    fmt_center_date = workbook.add_format(
        {
            "border": 1,
            "align": "center",
            "valign": "top",
            "text_wrap": True,
            "font_name": _FONT_NAME,
            "font_size": _FONT_SIZE,
            "num_format": "yyyy-mm-dd",
        }
    )
    fmt_std_part = workbook.add_format({"font_name": _FONT_NAME, "font_size": _FONT_SIZE})
    fmt_bold_part = workbook.add_format({"bold": True, "font_name": _FONT_NAME, "font_size": _FONT_SIZE})
    fmt_bold_cell_left = workbook.add_format(
        {
            "bold": True,
            "border": 1,
            "align": "left",
            "valign": "top",
            "text_wrap": True,
            "font_name": _FONT_NAME,
            "font_size": _FONT_SIZE,
        }
    )
    fmt_bold_cell_center = workbook.add_format(
        {
            "bold": True,
            "border": 1,
            "align": "center",
            "valign": "top",
            "text_wrap": True,
            "font_name": _FONT_NAME,
            "font_size": _FONT_SIZE,
        }
    )

    for c_idx, cell in enumerate(data_rows[0]):
        header = cell.processed_text.lower()
        width = 30
        if "position" in header:
            width = 22
        elif "name" in header:
            width = 18
        elif "total" in header:
            width = 12
        elif "regional" in header:
            width = 20
        elif "experience" in header:
            width = 65
        elif "education" in header:
            width = 45
        elif "role" in header:
            width = 40
        elif "qualification" in header:
            width = 60
        worksheet.set_column(c_idx, c_idx, width)

    for r_idx, row in enumerate(data_rows):
        for c_idx, cell in enumerate(row):
            if cell.is_header:
                worksheet.write(r_idx, c_idx, cell.processed_text, fmt_header)
                continue
            val = cell.processed_text
            if not val:
                worksheet.write_blank(r_idx, c_idx, "", fmt_std)
                continue
            raw = cell.raw_value
            if isinstance(raw, str) and raw.startswith("="):
                worksheet.write_formula(r_idx, c_idx, raw, fmt_center if cell.is_center else fmt_std)
                continue
            keywords = cell.keywords
            base_fmt = fmt_center if cell.is_center else fmt_std
            # Numbers and dates keep their type. Previously every cell
            # fell through to the text path below, so a numeric 2.5
            # came back as the string "2.5" and any downstream SUM
            # silently ignored it. Keywords only apply to text cells.
            if isinstance(raw, bool):
                worksheet.write_boolean(r_idx, c_idx, raw, base_fmt)
                continue
            if isinstance(raw, (int, float)) and math.isfinite(raw):
                worksheet.write_number(r_idx, c_idx, raw, base_fmt)
                continue
            date_fmt = fmt_center_date if cell.is_center else fmt_std_date
            if isinstance(raw, datetime.datetime):
                worksheet.write_datetime(r_idx, c_idx, raw, date_fmt)
                continue
            if isinstance(raw, datetime.date):
                worksheet.write_datetime(r_idx, c_idx, datetime.datetime(raw.year, raw.month, raw.day), date_fmt)
                continue
            if keywords and val:
                sorted_keywords = sorted(set(keywords), key=len, reverse=True)
                pattern = r"(" + "|".join(re.escape(k) for k in sorted_keywords) + r")"
                parts = re.split(pattern, val)
                rich_parts: list[Any] = []
                has_bold = False
                for part in parts:
                    if not part:
                        continue
                    if any(part.lower() == k.lower() for k in keywords):
                        rich_parts.append(fmt_bold_part)
                        rich_parts.append(part)
                        has_bold = True
                    else:
                        rich_parts.append(fmt_std_part)
                        rich_parts.append(part)
                string_frags = [p for p in rich_parts if isinstance(p, str)]
                if has_bold and len(string_frags) > 1:
                    try:
                        worksheet.write_rich_string(r_idx, c_idx, *rich_parts, base_fmt)
                    except Exception:
                        logger.warning(
                            "Rich-text write failed at row %s col %s; writing plain text",
                            r_idx,
                            c_idx,
                        )
                        worksheet.write(r_idx, c_idx, val, base_fmt)
                elif has_bold and len(string_frags) == 1:
                    worksheet.write(
                        r_idx,
                        c_idx,
                        val,
                        fmt_bold_cell_center if cell.is_center else fmt_bold_cell_left,
                    )
                else:
                    worksheet.write(r_idx, c_idx, val, base_fmt)
            else:
                worksheet.write(r_idx, c_idx, val, base_fmt)

    worksheet.set_landscape()
    worksheet.fit_to_pages(1, 0)
    workbook.close()
    return out_path


def format_matrix(
    source_path: str | Path,
    output_path: str | Path,
    docx_path: str | Path | None = None,
    sheet_name: str | None = None,
) -> Path:
    """Format one matrix workbook, optionally applying bold phrases."""
    bold_map = get_bold_map(docx_path) if docx_path else {}
    data_rows = extract_matrix_data(source_path, bold_map, sheet_name=sheet_name)
    return write_matrix(data_rows, Path(output_path))
