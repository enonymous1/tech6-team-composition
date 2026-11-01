"""Generate a FORM TECH-6 Excel workbook from a validated config.

The workbook is always built on a fresh sheet. v1 loaded a template
workbook, cleared its values in place, and wrote over it; stale
conditional formatting and empty styled rows from the template survived
into every output, and some configs overwrote the template file itself.
Nothing in the template contributed formatting the generator did not
already apply, so v2 does not load one.

Layout (columns shift with the number of deliverables):

    A: N°   B: Name   C: Position   D: [Home]/[Field]
    E..: one column per deliverable label
    Home / Field / Total: summed by live Excel formulas

Totals use plain ``SUM`` formulas. v1 stored bracketed strings such as
``[2 month]`` in cells and parsed them back with a
``SUBSTITUTE``/``VALUE``/``SEARCH`` formula. v2 parses the leading
number when the config is read (see :func:`parse_input_value`), so the
cell holds a number and the formula stays readable. Non-numeric text
is preserved as text; ``SUM`` ignores it.
"""

from __future__ import annotations

import logging
import re
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.worksheet import Worksheet

from .config import ConfigError, Tech6Config, load_config

logger = logging.getLogger(__name__)

_NUMBER_RE = re.compile(r"[-+]?\d+(?:\.\d+)?")
_FONT_NAME = "Gill Sans MT"
_FONT_SIZE = 10


def parse_input_value(value: Any) -> float | str | None:
    """Convert a config input cell to a number, text, or blank.

    Numbers pass through as floats. Strings yield their leading number
    (``"[2 month]"`` -> 2.0, ``"[0.5 m]"`` -> 0.5, ``"1.0"`` -> 1.0) so
    bracketed proposal shorthand still works. A string with no number
    is kept as text, and empty/blank values become ``None``.
    """
    if value is None:
        return None
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().strip("[]").strip()
    if not text:
        return None
    # Thousands separators only count when the whole string is numeric
    # ("1,000" -> 1000.0); "1,000 apples" still yields its leading 1.0
    # below, matching the existing leading-number behaviour.
    if re.fullmatch(r"[-+]?[\d,]+(?:\.\d+)?", text):
        text = text.replace(",", "")
    match = _NUMBER_RE.match(text)
    if match:
        return float(match.group(0))
    return text


def extract_threshold(text: str | None) -> float | None:
    """Extract the required total from text like ``"2.0 (1.0 Field)"``.

    Only a LEADING number counts: ``"2.0 (1.0 Field / 1.0 Remote)"``
    -> 2.0, and the parenthetical split is descriptive prose, not data.
    A string whose first number appears later ("Team 2") yields None,
    so it cannot silently become a requirement of 2.0.
    """
    if not text:
        return None
    match = _NUMBER_RE.match(str(text).strip())
    if match:
        return float(match.group(0))
    return None


class Tech6Generator:
    """Build a TECH-6 workbook from a :class:`Tech6Config`."""

    def __init__(self, config: Tech6Config):
        self.config = config
        self.project_info = config.project_info
        self.settings = config.settings
        self.d_labels = list(config.settings.deliverable_labels)

        self.col_n = 1
        self.col_name = 2
        self.col_pos = 3
        self.col_type = 4
        self.col_d_start = 5
        self.col_d_end = self.col_d_start + len(self.d_labels) - 1
        self.col_home = self.col_d_end + 1
        self.col_field = self.col_d_end + 2
        self.col_total = self.col_d_end + 3
        self.last_col = self.col_total

        self.font_std = Font(name=_FONT_NAME, size=_FONT_SIZE)
        self.font_bold = Font(name=_FONT_NAME, size=_FONT_SIZE, bold=True)
        self.font_title = Font(name=_FONT_NAME, size=_FONT_SIZE, bold=True)
        self.fill_header = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")
        self.fill_hatch = PatternFill(patternType="lightGrid", fgColor="D3D3D3", bgColor="FFFFFF")
        self.fill_section = PatternFill(start_color="DEEBF7", end_color="DEEBF7", fill_type="solid")
        self.border_thin = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )
        self.locked = Protection(locked=True)
        self.unlocked = Protection(locked=False)

    @classmethod
    def from_config_file(cls, path: str | Path) -> Tech6Generator:
        return cls(load_config(path))

    def resolve_output_path(
        self, output_dir: str | Path | None = None, output_filename: str | Path | None = None
    ) -> Path:
        """Resolve where the workbook will be written.

        Precedence: explicit ``output_dir`` argument, then the config's
        ``output_directory``, then ``./output``. The output path is
        refused if it would overwrite the config file itself. The
        filename override is applied BEFORE that check: applying it
        afterwards used to let ``--output-name`` point at the config
        file and overwrite it with a workbook.
        """
        filename = output_filename if output_filename is not None else self.project_info.output_filename
        if Path(filename).name != str(filename):
            raise ConfigError(f"output_filename must be a plain file name, got {filename!r}")
        if output_dir is not None:
            out_path = Path(output_dir) / filename
        elif self.config.output_directory is not None:
            out_path = self.config.output_directory / filename
        else:
            out_path = Path("output") / filename
        if self.config.source_path is not None and out_path.resolve() == self.config.source_path.resolve():
            raise ConfigError(f"Refusing to write output over the config file: {out_path}")
        return out_path

    def generate(self, output_dir: str | Path | None = None, output_filename: str | None = None) -> Path:
        """Build the workbook and save it, returning the output path."""
        out_path = self.resolve_output_path(output_dir, output_filename)
        workbook = self.build_workbook()
        out_path.parent.mkdir(parents=True, exist_ok=True)
        try:
            workbook.save(out_path)
        except OSError as exc:
            raise ConfigError(f"Could not save to {out_path}. Close the file if it is open in Excel: {exc}") from exc
        logger.info("Saved: %s", out_path)
        return out_path

    def build_workbook(self) -> Workbook:
        """Build the workbook in memory without writing a file."""
        workbook = openpyxl.Workbook()
        ws = workbook.active
        assert ws is not None
        ws.title = self.project_info.sheet_name

        row_idx = self._add_titles(ws)
        row_idx = self._add_headers(ws, row_idx)
        content_end_row = self._add_content(ws, row_idx)
        self._add_legend(ws, content_end_row)
        self._finalize(ws, row_idx)
        self._normalize_fonts(ws)
        return workbook

    def _add_titles(self, ws: Worksheet) -> int:
        row_idx = 1
        for i, title in enumerate(self.project_info.titles):
            cell = ws.cell(row=row_idx, column=1, value=title)
            cell.font = self.font_title if i > 0 else self.font_std
            cell.protection = self.locked
            if i > 0:
                cell.alignment = Alignment(horizontal="center")
                ws.merge_cells(
                    start_row=row_idx,
                    start_column=1,
                    end_row=row_idx,
                    end_column=self.last_col,
                )
            row_idx += 1
        return row_idx + 1

    def _add_headers(self, ws: Worksheet, row_idx: int) -> int:
        h_start = row_idx
        ws.cell(row=h_start, column=self.col_n, value="N°").font = self.font_bold
        ws.merge_cells(
            start_row=h_start,
            start_column=self.col_n,
            end_row=h_start + 1,
            end_column=self.col_n,
        )
        ws.cell(row=h_start, column=self.col_name, value="Name").font = self.font_bold
        ws.merge_cells(
            start_row=h_start,
            start_column=self.col_name,
            end_row=h_start + 1,
            end_column=self.col_name,
        )

        c_exp = ws.cell(
            row=h_start,
            column=self.col_pos,
            value="Expert's input (in person/month) per each Deliverable* (listed in TECH-5)",
        )
        c_exp.font = self.font_bold
        c_exp.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
        ws.merge_cells(
            start_row=h_start,
            start_column=self.col_pos,
            end_row=h_start,
            end_column=self.col_d_end,
        )

        ws.cell(row=h_start + 1, column=self.col_pos, value="Position").font = self.font_bold
        ws.cell(row=h_start + 1, column=self.col_type, value="").font = self.font_bold
        for i, label in enumerate(self.d_labels):
            ws.cell(row=h_start + 1, column=self.col_d_start + i, value=label).font = self.font_bold

        c_tot = ws.cell(row=h_start, column=self.col_home, value="Total time-input\n(in Months)")
        c_tot.font = self.font_bold
        c_tot.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.merge_cells(
            start_row=h_start,
            start_column=self.col_home,
            end_row=h_start,
            end_column=self.col_total,
        )

        ws.cell(row=h_start + 1, column=self.col_home, value="Home").font = self.font_bold
        ws.cell(row=h_start + 1, column=self.col_field, value="Field").font = self.font_bold
        ws.cell(row=h_start + 1, column=self.col_total, value="Total").font = self.font_bold
        ws.row_dimensions[h_start].height = 25

        for r in range(h_start, h_start + 2):
            for c in range(1, self.last_col + 1):
                cell = ws.cell(row=r, column=c)
                cell.fill = self.fill_header
                cell.border = self.border_thin
                cell.protection = self.locked
                is_main_header = r == h_start and c == self.col_pos
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=not is_main_header)
        return h_start + 2

    def _sum_formula(self, row_idx: int) -> str:
        start = get_column_letter(self.col_d_start)
        end = get_column_letter(self.col_d_end)
        return f"=SUM({start}{row_idx}:{end}{row_idx})"

    def _apply_validation(self, ws: Worksheet, cells: list[str], threshold: float, operator: str = "lessThan") -> None:
        """Flag cells red when they fail a minimum-input requirement.

        All validation in this generator is a minimum check: a cell is
        flagged when its total is *below* the required amount. v1's
        section-level check used ``greaterThan`` (flagging totals above
        the threshold), the opposite of the per-expert and grand-total
        checks; v2 uses ``lessThan`` everywhere.
        """
        red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        red_font = Font(name=_FONT_NAME, color="9C0006", bold=True)
        rule = CellIsRule(
            operator=operator,
            formula=[f"{threshold:g}"],
            stopIfTrue=True,
            fill=red_fill,
            font=red_font,
        )
        for coord in cells:
            ws.conditional_formatting.add(coord, rule)

    def _add_content(self, ws: Worksheet, start_row: int) -> int:
        row_idx = start_row
        subtotal_rows: list[int] = []

        for section in self.config.sections:
            c_sec = ws.cell(row=row_idx, column=1, value=section.header)
            c_sec.font = self.font_bold
            c_sec.fill = self.fill_section
            c_sec.border = self.border_thin
            c_sec.protection = self.locked
            ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=self.last_col)
            row_idx += 1

            home_cells: list[str] = []
            field_cells: list[str] = []
            key_expert_total_cells: list[str] = []

            for expert in section.experts:
                curr_start_r = row_idx
                expert_threshold = extract_threshold(expert.total_required_input)
                if expert_threshold is not None and expert_threshold < 0:
                    raise ConfigError(
                        f"Expert {expert.n} total_required_input must not be negative, "
                        f"got {expert.total_required_input!r}"
                    )

                for r in range(row_idx, row_idx + 2):
                    for c in range(1, self.last_col + 1):
                        cell = ws.cell(row=r, column=c)
                        cell.font = self.font_std
                        cell.border = self.border_thin

                # Row 1: Home.
                ws.cell(row=row_idx, column=self.col_n, value=expert.n).protection = self.unlocked
                ws.cell(row=row_idx, column=self.col_name, value=expert.name).protection = self.unlocked
                ws.cell(row=row_idx, column=self.col_pos, value=expert.position).protection = self.unlocked
                c_type_h = ws.cell(row=row_idx, column=self.col_type, value="[Home]")
                c_type_h.protection = self.locked
                c_type_h.alignment = Alignment(horizontal="center")

                h_in_vals = list(expert.h_in) + [""] * (len(self.d_labels) - len(expert.h_in))
                for i, value in enumerate(h_in_vals):
                    c_in = ws.cell(row=row_idx, column=self.col_d_start + i, value=parse_input_value(value))
                    c_in.alignment = Alignment(horizontal="center")
                    c_in.protection = self.unlocked
                    c_in.font = self.font_std

                c_home_sum = ws.cell(row=row_idx, column=self.col_home)
                c_home_sum.value = self._sum_formula(row_idx)
                c_home_sum.protection = self.locked
                c_home_sum.font = self.font_std
                c_home_sum.alignment = Alignment(horizontal="center")
                home_cells.append(c_home_sum.coordinate)
                ws.cell(row=row_idx, column=self.col_field).fill = self.fill_hatch
                row_idx += 1

                # Row 2: Field.
                c_type_f = ws.cell(row=row_idx, column=self.col_type, value="[Field]")
                c_type_f.protection = self.locked
                c_type_f.alignment = Alignment(horizontal="center")

                f_in_vals = list(expert.f_in) + [""] * (len(self.d_labels) - len(expert.f_in))
                for i, value in enumerate(f_in_vals):
                    c_in = ws.cell(row=row_idx, column=self.col_d_start + i, value=parse_input_value(value))
                    c_in.alignment = Alignment(horizontal="center")
                    c_in.protection = self.unlocked
                    c_in.font = self.font_std

                ws.cell(row=row_idx, column=self.col_home).fill = self.fill_hatch
                c_field_sum = ws.cell(row=row_idx, column=self.col_field)
                c_field_sum.value = self._sum_formula(row_idx)
                c_field_sum.protection = self.locked
                c_field_sum.font = self.font_std
                c_field_sum.alignment = Alignment(horizontal="center")
                field_cells.append(c_field_sum.coordinate)

                total_formula = (
                    f"={get_column_letter(self.col_home)}{row_idx - 1} + {get_column_letter(self.col_field)}{row_idx}"
                )
                c_total = ws.cell(row=row_idx - 1, column=self.col_total, value=total_formula)
                ws.merge_cells(
                    start_row=row_idx - 1,
                    start_column=self.col_total,
                    end_row=row_idx,
                    end_column=self.col_total,
                )
                c_total.alignment = Alignment(horizontal="center", vertical="center")
                c_total.protection = self.locked
                c_total.font = self.font_std

                if expert_threshold is not None:
                    self._apply_validation(ws, [c_total.coordinate], expert_threshold)
                elif section.is_key_experts:
                    key_expert_total_cells.append(c_total.coordinate)

                for c_idx in (self.col_n, self.col_name, self.col_pos):
                    ws.merge_cells(start_row=curr_start_r, start_column=c_idx, end_row=row_idx, end_column=c_idx)
                    ws.cell(row=curr_start_r, column=c_idx).alignment = Alignment(
                        horizontal="center", vertical="center", wrap_text=True
                    )
                row_idx += 1

            # Subtotal row.
            sub_row = row_idx
            subtotal_rows.append(sub_row)
            ws.merge_cells(start_row=sub_row, start_column=1, end_row=sub_row, end_column=self.col_d_end)
            c_sub = ws.cell(row=sub_row, column=1, value="Subtotal")
            c_sub.font = self.font_bold
            c_sub.alignment = Alignment(horizontal="right")

            h_sum_f = f"=SUM({','.join(home_cells)})" if home_cells else 0
            f_sum_f = f"=SUM({','.join(field_cells)})" if field_cells else 0
            t_sum_f = f"={get_column_letter(self.col_home)}{sub_row}+{get_column_letter(self.col_field)}{sub_row}"
            ws.cell(row=sub_row, column=self.col_home, value=h_sum_f).font = self.font_bold
            ws.cell(row=sub_row, column=self.col_field, value=f_sum_f).font = self.font_bold
            ws.cell(row=sub_row, column=self.col_total, value=t_sum_f).font = self.font_bold

            for c in range(1, self.last_col + 1):
                cell = ws.cell(row=sub_row, column=c)
                cell.border = self.border_thin
                cell.protection = self.locked
                if c >= self.col_home:
                    cell.alignment = Alignment(horizontal="center")

            if section.is_key_experts and key_expert_total_cells:
                self._apply_validation(ws, key_expert_total_cells, self.settings.validation_threshold)
            row_idx += 1

        self._add_grand_total(ws, row_idx, subtotal_rows)
        return row_idx

    def _add_grand_total(self, ws: Worksheet, row_idx: int, subtotal_rows: list[int]) -> None:
        ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=self.col_d_end)
        c_tot = ws.cell(row=row_idx, column=1, value="Total")
        c_tot.font = self.font_bold
        c_tot.alignment = Alignment(horizontal="right")

        def multi_sum(col: int) -> str | int:
            if not subtotal_rows:
                return 0
            return "=" + "+".join(get_column_letter(col) + str(r) for r in subtotal_rows)

        c_home = ws.cell(row=row_idx, column=self.col_home, value=multi_sum(self.col_home))
        c_field = ws.cell(row=row_idx, column=self.col_field, value=multi_sum(self.col_field))
        c_total = ws.cell(row=row_idx, column=self.col_total, value=multi_sum(self.col_total))
        c_home.font = self.font_bold
        c_field.font = self.font_bold
        c_total.font = self.font_bold

        total_min = self.settings.total_minimum_months
        field_min = self.settings.field_minimum_months
        home_min = self.settings.home_minimum_months
        if home_min is None and total_min is not None and field_min is not None:
            home_min = total_min - field_min

        if total_min is not None:
            self._apply_validation(ws, [c_total.coordinate], total_min)
        if field_min is not None:
            self._apply_validation(ws, [c_field.coordinate], field_min)
        if home_min is not None:
            self._apply_validation(ws, [c_home.coordinate], home_min)

        for c in range(1, self.last_col + 1):
            cell = ws.cell(row=row_idx, column=c)
            cell.border = self.border_thin
            cell.protection = self.locked
            if c >= self.col_home:
                cell.alignment = Alignment(horizontal="center")

    def _add_legend(self, ws: Worksheet, row_idx: int) -> int:
        row_idx += 2
        ref_text = "* Full descriptions of Deliverables are provided in TECH-5 [PG. #]"
        ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=self.last_col)
        c_ref = ws.cell(row=row_idx, column=1, value=ref_text)
        c_ref.font = Font(name=_FONT_NAME, size=_FONT_SIZE, bold=True, italic=True)
        c_ref.alignment = Alignment(horizontal="left")
        c_ref.border = Border()
        row_idx += 2

        for deliverable in self.config.deliverables:
            ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=self.last_col)
            c_title = ws.cell(row=row_idx, column=1, value=f"{deliverable.id}. {deliverable.title}")
            c_title.font = self.font_bold
            c_title.alignment = Alignment(horizontal="left")
            c_title.border = Border()
            row_idx += 1

            ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=self.last_col)
            c_subtitle = ws.cell(row=row_idx, column=1, value=deliverable.subtitle)
            c_subtitle.font = Font(name=_FONT_NAME, size=_FONT_SIZE, italic=True)
            c_subtitle.alignment = Alignment(horizontal="left")
            c_subtitle.border = Border()
            row_idx += 1

            ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=self.last_col)
            c_subtext = ws.cell(row=row_idx, column=1, value=f'"{deliverable.subtext}"')
            c_subtext.font = Font(name=_FONT_NAME, size=_FONT_SIZE)
            c_subtext.alignment = Alignment(horizontal="left")
            c_subtext.border = Border()
            row_idx += 2
        return row_idx

    def _finalize(self, ws: Worksheet, first_content_row: int) -> None:
        ws.column_dimensions["A"].width = 8
        ws.column_dimensions["B"].width = 25
        ws.column_dimensions["C"].width = 30
        ws.column_dimensions["D"].width = 10
        for i in range(len(self.d_labels)):
            ws.column_dimensions[get_column_letter(self.col_d_start + i)].width = 12
        ws.column_dimensions[get_column_letter(self.col_home)].width = 9
        ws.column_dimensions[get_column_letter(self.col_field)].width = 9
        ws.column_dimensions[get_column_letter(self.col_total)].width = 9

        # Sheet protection ON: the generator marks input cells unlocked
        # and formula/label cells locked, and those flags only mean
        # anything with protection enabled. Without a password, so a
        # proposal author can still unprotect it deliberately.
        ws.protection.sheet = True
        ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        # Freeze below the header rows so the table stays readable.
        ws.freeze_panes = f"A{first_content_row}"

    def _normalize_fonts(self, ws: Worksheet) -> None:
        """Ensure every written cell uses the house font at size 10."""
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
            for cell in row:
                if not cell.font or cell.font.name != _FONT_NAME:
                    cell.font = Font(
                        name=_FONT_NAME,
                        size=_FONT_SIZE,
                        bold=cell.font.bold if cell.font else False,
                        italic=cell.font.italic if cell.font else False,
                        color=cell.font.color if cell.font else None,
                    )
