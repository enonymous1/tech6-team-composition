"""Tests for the TECH-6 workbook generator."""

import openpyxl
import pytest

from tech6_team_composition.config import ConfigError
from tech6_team_composition.generator import (
    Tech6Generator,
    extract_threshold,
    parse_input_value,
)


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        (None, None),
        ("", None),
        ("   ", None),
        (0.5, 0.5),
        (2, 2.0),
        ("0.5", 0.5),
        ("[1.0]", 1.0),
        ("[2 month]", 2.0),
        ("[0.5 m]", 0.5),
        ("1.0 Field", 1.0),
        ("TBD", "TBD"),
        ("[TBD]", "TBD"),
    ],
)
def test_parse_input_value(value, expected):
    assert parse_input_value(value) == expected


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        (None, None),
        ("", None),
        ("2.0 (1.0 Field / 1.0 Remote)", 2.0),
        ("1.5 (0.5 Field / 1.0 Remote)", 1.5),
        ("2.0", 2.0),
        ("no numbers here", None),
    ],
)
def test_extract_threshold(text, expected):
    assert extract_threshold(text) == expected


def _workbook_for(examples_dir, name):
    generator = Tech6Generator.from_config_file(examples_dir / name)
    return generator.build_workbook(), generator


def test_validation_example_layout(examples_dir):
    workbook, generator = _workbook_for(examples_dir, "validation.json")
    ws = workbook.active
    assert ws is not None
    assert ws.title == "Validation Test"
    # Titles, blank row, 2 header rows, section header, 10 expert rows,
    # subtotal, grand total, then the legend. No ghost rows from a template.
    assert ws["A1"].value.startswith("Section 3")
    assert ws["A5"].value == "N°"
    assert ws["A7"].value == "KEY EXPERTS"
    assert ws["B8"].value == "Mr. Team Leader"
    assert ws["A18"].value == "Subtotal"
    assert ws["A19"].value == "Total"
    assert ws["A21"].value.startswith("* Full descriptions")
    assert ws["A23"].value == "D-1. Inception Report"
    assert ws.max_row == 37
    assert ws.max_column == generator.last_col == 11


def test_validation_example_formulas(examples_dir):
    workbook, generator = _workbook_for(examples_dir, "validation.json")
    ws = workbook.active
    assert ws is not None
    # Home/field sums are plain SUMs over the deliverable columns.
    assert ws["I8"].value == "=SUM(E8:H8)"
    assert ws["J9"].value == "=SUM(E9:H9)"
    # Total merges home + field.
    assert ws["K8"].value == "=I8 + J9"
    # Subtotal sums the per-expert sums; grand total sums subtotals.
    assert ws["I18"].value == "=SUM(I8,I10,I12,I14,I16)"
    assert ws["J18"].value == "=SUM(J9,J11,J13,J15,J17)"
    assert ws["I19"].value == "=I18"
    assert ws["K19"].value == "=K18"
    # Bracketed shorthand became numbers in cells.
    assert ws["E8"].value == 0.5
    assert ws["F9"].value == 0.5


def test_full_example_bracketed_inputs_become_numbers(examples_dir):
    workbook, _ = _workbook_for(examples_dir, "full.json")
    ws = workbook.active
    assert ws is not None
    assert ws["E8"].value == 2.0
    assert ws["F8"].value == 1.0
    assert ws["E9"].value == 0.5
    assert ws["F9"].value == 2.5
    # Seven deliverables: home is L, field is M, total is N.
    assert ws["L8"].value == "=SUM(E8:K8)"
    assert ws["M9"].value == "=SUM(E9:K9)"
    assert ws["N8"].value == "=L8 + M9"


def test_minimal_example_dimensions(examples_dir):
    workbook, generator = _workbook_for(examples_dir, "minimal.json")
    ws = workbook.active
    assert ws is not None
    assert ws.max_column == generator.last_col == 8
    assert ws["E6"].value == "D-1"
    assert ws["H8"].value == "=F8 + G9"


def test_no_conditional_formatting_beyond_last_column(examples_dir):
    workbook, generator = _workbook_for(examples_dir, "validation.json")
    ws = workbook.active
    assert ws is not None
    import re

    from openpyxl.utils import column_index_from_string

    for cf_range in ws.conditional_formatting._cf_rules:
        for cell_ref in str(cf_range.sqref).split():
            col_letters = re.match(r"[A-Z]+", cell_ref)
            assert col_letters is not None
            assert column_index_from_string(col_letters.group(0)) <= generator.last_col


def _rules_by_range(ws):
    return {str(key.sqref): value for key, value in ws.conditional_formatting._cf_rules.items()}


def test_validation_rules_are_minimum_checks(examples_dir):
    workbook, _ = _workbook_for(examples_dir, "validation.json")
    ws = workbook.active
    assert ws is not None
    rules = _rules_by_range(ws)
    # Per-expert totals use lessThan against their required input.
    assert rules["K8"][0].operator == "lessThan"
    assert rules["K8"][0].formula == ["2"]
    assert rules["K14"][0].operator == "lessThan"
    assert rules["K14"][0].formula == ["1.5"]
    # Grand totals use lessThan against the stated minimums.
    assert rules["K19"][0].formula == ["13.5"]
    assert rules["J19"][0].formula == ["7.5"]
    assert rules["I19"][0].formula == ["6"]
    # No duplicated rules on any cell.
    for cf_range, rule_list in rules.items():
        assert len(rule_list) == 1, cf_range


def test_section_threshold_is_a_minimum_check(examples_dir):
    # Basic example experts have no individual requirement, so the
    # section threshold applies, as a lessThan (minimum) check.
    workbook, _ = _workbook_for(examples_dir, "basic.json")
    ws = workbook.active
    assert ws is not None
    rules = _rules_by_range(ws)
    # Seven deliverables: total column is N.
    assert rules["N8"][0].operator == "lessThan"
    assert rules["N8"][0].formula == ["6"]


def test_repeated_builds_do_not_accumulate_rules(examples_dir):
    generator = Tech6Generator.from_config_file(examples_dir / "validation.json")
    first = generator.build_workbook().active
    second = generator.build_workbook().active
    assert first is not None and second is not None
    assert len(first.conditional_formatting._cf_rules) == len(second.conditional_formatting._cf_rules)


def test_fonts_are_uniform(examples_dir):
    workbook, _ = _workbook_for(examples_dir, "validation.json")
    ws = workbook.active
    assert ws is not None
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            if cell.value is not None:
                assert cell.font.name == "Gill Sans MT"
                assert cell.font.size == 10


def test_generate_writes_file_and_respects_output_dir(examples_dir, tmp_path):
    generator = Tech6Generator.from_config_file(examples_dir / "minimal.json")
    out = generator.generate(output_dir=tmp_path / "out")
    assert out.exists()
    assert out.name == "tech6_minimal_example.xlsx"
    workbook = openpyxl.load_workbook(out)
    assert workbook.active is not None
    assert workbook.active.title == "Team Input"


def test_generate_output_name_override(examples_dir, tmp_path):
    generator = Tech6Generator.from_config_file(examples_dir / "minimal.json")
    out = generator.generate(output_dir=tmp_path, output_filename="custom.xlsx")
    assert out.name == "custom.xlsx"


def test_generate_twice_is_stable(examples_dir, tmp_path):
    generator = Tech6Generator.from_config_file(examples_dir / "validation.json")
    first = generator.generate(output_dir=tmp_path)
    second = generator.generate(output_dir=tmp_path)
    assert first == second
    ws = openpyxl.load_workbook(second).active
    assert ws is not None
    rules = ws.conditional_formatting._cf_rules
    assert all(len(v) == 1 for v in rules.values())


def test_generate_refuses_to_overwrite_config(examples_dir, tmp_path):
    import json
    import shutil

    config_path = tmp_path / "cfg.json"
    shutil.copy(examples_dir / "minimal.json", config_path)
    data = json.loads(config_path.read_text(encoding="utf-8"))
    data["project_info"]["output_filename"] = "cfg.json"
    config_path.write_text(json.dumps(data), encoding="utf-8")
    generator = Tech6Generator.from_config_file(config_path)
    with pytest.raises(ConfigError, match="Refusing to write output over the config"):
        generator.generate(output_dir=tmp_path)


def test_missing_config_file():
    with pytest.raises(ConfigError, match="not found"):
        Tech6Generator.from_config_file("no-such-config.json")


def test_page_setup_and_freeze(examples_dir):
    workbook, _ = _workbook_for(examples_dir, "validation.json")
    ws = workbook.active
    assert ws is not None
    assert ws.page_setup.orientation == "landscape"
    assert ws.sheet_properties.pageSetUpPr.fitToPage is True
    assert ws.page_setup.fitToWidth == 1
    assert ws.freeze_panes is not None
