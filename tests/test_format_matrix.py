"""Tests for the qualifications matrix formatter."""

import openpyxl
import pytest
from docx import Document

from tech6_team_composition.format_matrix import (
    extract_matrix_data,
    format_matrix,
    get_bold_map,
    write_matrix,
)


def _matrix_workbook(tmp_path, name="matrix.xlsx"):
    path = tmp_path / name
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(
        ["Position", "Name", "Total Input", "Education & Certification", "Experience Highlights", "Regional Exp."]
    )
    ws.append(
        [
            "K-1: Team Leader",
            "Jane Q. Expert",
            "2.0 Months",
            "MSc in Project Management <br>• PMP Certified",
            "Led a fictional programme <br>• Built a fictional portal",
            "Yes",
        ]
    )
    ws.append(
        ["K-2: Architect", "Alex Example", "2.0 Months", "PhD <br>• Specialist", "Designed a fictional platform", "Yes"]
    )
    wb.save(path)
    return path


def _styles_docx(tmp_path):
    path = tmp_path / "styles.docx"
    doc = Document()
    table = doc.add_table(rows=1, cols=6)
    for i, header in enumerate(["Position", "Name", "Total", "Education", "Experience", "Regional"]):
        table.rows[0].cells[i].text = header
    row = table.add_row().cells
    row[1].text = "Jane Q. Expert"
    run = row[3].paragraphs[0].add_run("MSc in Project Management")
    run.bold = True
    doc.save(path)
    return path


def test_get_bold_map(tmp_path):
    docx_path = _styles_docx(tmp_path)
    bold_map = get_bold_map(docx_path)
    assert bold_map[("Jane Q. Expert", 3)] == ["MSc in Project Management"]


def test_get_bold_map_missing_file(tmp_path):
    with pytest.raises(FileNotFoundError):
        get_bold_map(tmp_path / "missing.docx")


def test_get_bold_map_no_tables(tmp_path):
    path = tmp_path / "empty.docx"
    Document().save(path)
    assert get_bold_map(path) == {}


def test_extract_matrix_data(tmp_path):
    source = _matrix_workbook(tmp_path)
    rows = extract_matrix_data(source, {("Jane Q. Expert", 3): ["MSc in Project Management"]})
    assert len(rows) == 3
    header, first, _ = rows
    assert header[0].is_header
    assert header[3].is_bullet  # Education column.
    assert header[1].is_center  # Name column.
    assert first[3].keywords == ("MSc in Project Management",)
    assert first[3].processed_text.startswith("• MSc in Project Management")
    assert "<br>" not in first[3].processed_text


def test_extract_missing_file(tmp_path):
    with pytest.raises(FileNotFoundError):
        extract_matrix_data(tmp_path / "missing.xlsx", {})


def test_write_matrix_round_trip(tmp_path):
    source = _matrix_workbook(tmp_path)
    rows = extract_matrix_data(source, {})
    out = tmp_path / "formatted.xlsx"
    result = write_matrix(rows, out)
    assert result == out
    wb = openpyxl.load_workbook(out)
    ws = wb.active
    assert ws is not None
    assert ws["A1"].value == "Position"
    assert ws["B2"].value == "Jane Q. Expert"
    assert ws["D2"].value.startswith("• MSc")


def test_write_matrix_empty_rows():
    with pytest.raises(ValueError, match="No matrix rows"):
        write_matrix([], "out.xlsx")


def test_format_matrix_end_to_end(tmp_path, examples_dir):
    out = tmp_path / "qualifications_Formatted.xlsx"
    result = format_matrix(
        examples_dir / "qualifications_matrix.xlsx",
        out,
        docx_path=examples_dir / "qualifications_styles.docx",
    )
    assert result.exists()
    wb = openpyxl.load_workbook(result)
    ws = wb.active
    assert ws is not None
    assert ws.max_row == 4
    assert ws["B2"].value == "Jane Q. Expert"


def test_format_matrix_without_docx(tmp_path, examples_dir):
    out = tmp_path / "plain.xlsx"
    result = format_matrix(examples_dir / "qualifications_matrix.xlsx", out)
    assert result.exists()


def test_bundled_example_files_are_fictional(examples_dir):
    wb = openpyxl.load_workbook(examples_dir / "qualifications_matrix.xlsx")
    ws = wb.active
    assert ws is not None
    names = [ws.cell(row=r, column=2).value for r in range(2, ws.max_row + 1)]
    assert names == ["Jane Q. Expert", "Alex Example", "Sam Sample"]
